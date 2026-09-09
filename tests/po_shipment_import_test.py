"""Workbook and CLI behavior at the file/HTTP boundaries (no live Supabase)."""
import contextlib
import io
import os
from pathlib import Path
import sys
import tempfile
import unittest
from datetime import datetime
from unittest.mock import Mock, patch

from openpyxl import Workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import sync_po_shipment_data as importer

PROJECT = "b1fb89d6-5321-4a31-b5ba-9e82f51204a2"
PO_HEADERS = ["Purchase Order ID", "Purchase Order Item", "Item Description",
              "Product", "Product", "Net Price", "Ordered Quantity", "Created On"]
SHIP_HEADERS = ["Shipment #", "PO#", "ETA", "# Pcs", "NCR/OSD (X)", "Extra notes"]


class WorkbookImportTests(unittest.TestCase):
    def setUp(self):
        temporary = tempfile.TemporaryDirectory()
        self.addCleanup(temporary.cleanup)
        self.path = Path(temporary.name) / "snapshot.xlsx"
        self.workbook = Workbook()
        po = self.workbook.active
        po.title = "PO Parts Log"
        po.append(PO_HEADERS)
        po.append([20001.0, 10.0, "Valve", "primary", "alternate", "123.45", "2.5",
                   datetime(2026, 9, 9)])
        shipment = self.workbook.create_sheet("Shipment Log")
        shipment.append(SHIP_HEADERS)
        shipment.append([101.0, 20001.0, datetime(2026, 9, 10), "3", "X", "Ignore me"])
        guard = patch("requests.sessions.Session.request", side_effect=AssertionError("Live HTTP forbidden"))
        self.network = guard.start()
        self.addCleanup(guard.stop)

    def read(self):
        self.workbook.save(self.path)
        return importer.read_workbook(self.path)

    def run_cli(self, apply=False, environment=None, project=PROJECT):
        self.workbook.save(self.path)
        argv = ["--file", str(self.path), "--project-id", project] + (["--apply"] if apply else [])
        with patch.dict(os.environ, environment or {}, clear=True), contextlib.redirect_stdout(io.StringIO()), contextlib.redirect_stderr(io.StringIO()):
            return importer.main(argv)

    def test_normalizes_schema_fields_without_losing_duplicate_product_columns(self):
        purchase_orders, shipments = self.read()
        po, shipment = purchase_orders[0], shipments[0]
        self.assertEqual((po["purchase_order_id"], po["purchase_order_item"]), ("20001", "10"))
        self.assertEqual((po["product"], po["product_alt"]), ("primary", "alternate"))
        self.assertEqual((po["net_price"], po["ordered_quantity"]), (123.45, 2.5))
        self.assertEqual(po["created_on"], "2026-09-09")
        self.assertEqual((shipment["shipment_number"], shipment["po_number"]), ("101", "20001"))
        self.assertEqual((shipment["eta"], shipment["num_pieces"], shipment["ncr_osd"]), ("2026-09-10", 3, True))
        self.assertNotIn("Extra notes", shipment)
        self.assertFalse({"id", "project_id", "synced_at"}.intersection(po))

    def test_ignores_blank_rows_and_allows_distinct_lines_on_same_po(self):
        self.workbook["PO Parts Log"].append([None] * len(PO_HEADERS))
        self.workbook["PO Parts Log"].append([20001, 20, "Pipe"])
        self.workbook["Shipment Log"].append([None] * len(SHIP_HEADERS))
        purchase_orders, shipments = self.read()
        self.assertEqual((len(purchase_orders), len(shipments)), (2, 1))

    def test_omits_absent_optional_columns_to_preserve_existing_server_values(self):
        self.workbook["PO Parts Log"].delete_cols(4, len(PO_HEADERS) - 3)
        self.workbook["Shipment Log"].delete_cols(2, len(SHIP_HEADERS) - 1)
        purchase_orders, shipments = self.read()
        self.assertEqual(purchase_orders, [{"purchase_order_id": "20001", "purchase_order_item": "10", "item_description": "Valve"}])
        self.assertEqual(shipments, [{"shipment_number": "101"}])

    def test_rejects_uncached_formulas_in_mapped_fields(self):
        cases = [("PO Parts Log", "C2", '="Valve"'), ("PO Parts Log", "F2", "=100+23.45"),
                 ("Shipment Log", "C2", "=DATE(2026,9,10)"), ("Shipment Log", "D2", "=1+2")]
        for sheet, coordinate, formula in cases:
            with self.subTest(sheet=sheet, coordinate=coordinate):
                cell = self.workbook[sheet][coordinate]
                original, cell.value = cell.value, formula
                with self.assertRaises(ValueError):
                    self.read()
                cell.value = original

    def test_rejects_excel_error_cells_in_text_fields(self):
        self.workbook["PO Parts Log"]["C2"] = "#REF!"
        with self.assertRaises(ValueError):
            self.read()

    def test_corrupt_xlsx_returns_failure_without_traceback_or_network(self):
        self.path.write_bytes(b"This is not an XLSX ZIP archive")
        errors = io.StringIO()
        with patch.dict(os.environ, {}, clear=True), contextlib.redirect_stderr(errors), contextlib.redirect_stdout(io.StringIO()):
            result = importer.main(["--file", str(self.path), "--project-id", PROJECT])
        self.assertEqual(result, 1)
        self.assertNotIn("Traceback", errors.getvalue())
        self.assertTrue(errors.getvalue().strip())
        self.network.assert_not_called()

    def test_rejects_each_missing_required_header(self):
        for sheet, column in [("PO Parts Log", 1), ("PO Parts Log", 2), ("PO Parts Log", 3), ("Shipment Log", 1)]:
            with self.subTest(sheet=sheet, column=column):
                cell = self.workbook[sheet].cell(1, column)
                original, cell.value = cell.value, "Unrecognized heading"
                with self.assertRaises(ValueError):
                    self.read()
                cell.value = original

    def test_rejects_missing_or_empty_sheet(self):
        for sheet in ["PO Parts Log", "Shipment Log"]:
            with self.subTest(sheet=sheet):
                worksheet = self.workbook[sheet]
                row = [cell.value for cell in worksheet[2]]
                worksheet.delete_rows(2, worksheet.max_row)
                with self.assertRaises(ValueError):
                    self.read()
                worksheet.append(row)
        del self.workbook["Shipment Log"]
        with self.assertRaises(ValueError):
            self.read()

    def test_rejects_duplicate_identifiers_in_either_sheet(self):
        for sheet in ["PO Parts Log", "Shipment Log"]:
            with self.subTest(sheet=sheet):
                worksheet = self.workbook[sheet]
                worksheet.append([cell.value for cell in worksheet[2]])
                with self.assertRaises(ValueError):
                    self.read()
                worksheet.delete_rows(3)

    def test_rejects_supplied_invalid_dates_numbers_and_blank_identifiers(self):
        cases = [("PO Parts Log", "H2", "2026-02-30"), ("Shipment Log", "C2", "tomorrow"),
                 ("PO Parts Log", "F2", "NaN"), ("PO Parts Log", "G2", "many"),
                 ("Shipment Log", "D2", "1.5"), ("PO Parts Log", "A2", None),
                 ("Shipment Log", "A2", None)]
        for sheet, coordinate, invalid in cases:
            with self.subTest(sheet=sheet, coordinate=coordinate, invalid=invalid):
                cell = self.workbook[sheet][coordinate]
                original, cell.value = cell.value, invalid
                with self.assertRaises(ValueError):
                    self.read()
                cell.value = original

    def test_dry_run_needs_no_credentials_and_never_contacts_server(self):
        self.assertEqual(self.run_cli(), 0)
        self.network.assert_not_called()

    def test_apply_requires_service_credential_and_https(self):
        for env in [{}, {"SUPABASE_URL": "https://example.supabase.co", "SUPABASE_ANON_KEY": "anon"},
                    {"SUPABASE_URL": "http://example.supabase.co", "SUPABASE_SERVICE_ROLE_KEY": "secret"}]:
            with self.subTest(environment_keys=list(env)):
                self.assertEqual(self.run_cli(apply=True, environment=env), 1)
        self.network.assert_not_called()

    def test_apply_posts_both_validated_sheets_in_one_transaction(self):
        response = Mock(status_code=200, ok=True)
        response.json.return_value = {"purchase_orders": 1, "shipments": 1, "synced_at": "2026-09-09T12:00:00Z"}
        with patch.object(importer.requests, "post", return_value=response) as post:
            self.assertEqual(self.run_cli(True, {"SUPABASE_URL": "https://example.supabase.co", "SUPABASE_SERVICE_ROLE_KEY": "secret"}), 0)
        post.assert_called_once()
        args, kwargs = post.call_args
        self.assertEqual(args[0], "https://example.supabase.co/rest/v1/rpc/import_po_shipment_snapshot")
        purchase_orders, shipments = self.read()
        self.assertEqual(kwargs["json"], {"p_project_id": PROJECT, "p_purchase_orders": purchase_orders, "p_shipments": shipments})
        self.assertEqual(kwargs["headers"]["Authorization"], "Bearer secret")

    def test_invalid_second_sheet_prevents_any_write(self):
        self.workbook["Shipment Log"]["C2"] = "not a date"
        with patch.object(importer.requests, "post") as post:
            self.assertEqual(self.run_cli(True, {"SUPABASE_URL": "https://example.supabase.co", "SUPABASE_SERVICE_ROLE_KEY": "secret"}), 1)
        post.assert_not_called()

    def test_http_failure_returns_failure_without_retry_or_other_write(self):
        response = Mock(status_code=409, ok=False, text="Conflicting shipment identifier")
        response.raise_for_status.side_effect = importer.requests.HTTPError("409 Conflict")
        with patch.object(importer.requests, "post", return_value=response) as post:
            self.assertEqual(self.run_cli(True, {"SUPABASE_URL": "https://example.supabase.co", "SUPABASE_SERVICE_ROLE_KEY": "secret"}), 1)
        post.assert_called_once()


if __name__ == "__main__":
    unittest.main()
