"""Validate an MSR workbook, then optionally import both sheets atomically.

Default: validation only. Apply requires an explicit project, HTTPS Supabase URL,
and a server-only SUPABASE_SERVICE_ROLE_KEY. No rows are deleted.
"""
import argparse
import math
import os
import sys
from datetime import date, datetime, time
from pathlib import Path
from urllib.parse import urlsplit
from uuid import UUID
from zipfile import BadZipFile

import requests
from openpyxl import load_workbook

PO_FIELDS = {
    "Purchase Order ID": "purchase_order_id", "PO Description": "po_description",
    "Purchase Order Item": "purchase_order_item", "Item UUID": "item_uuid",
    "Created On": "created_on", "Item Last Change Date Time": "item_last_change_date_time",
    "Delivery Date From": "delivery_date_from", "Status": "status", "Item Status": "item_status",
    "Delivery Status": "delivery_status", "Scope": "scope", "PO LI": "po_li", "Shipment": "shipment",
    "Category": "category", "Sub Category": "sub_category", "Project Task": "project_task",
    "Supplier": "supplier", "Item Description": "item_description",
    "Item Remark for Supplier": "item_remark_for_supplier", "Supplier Part Number": "supplier_part_number",
    "Product": "product", "Product.1": "product_alt", "Manufacturer": "manufacturer",
    "Manufacturer Part Number": "manufacturer_part_number", "Base UoM": "base_uom",
    "Item Type": "item_type", "Ordered Quantity": "ordered_quantity",
    "Base Net Price Base Quantity Unit": "base_net_price_base_quantity_unit",
    "Net Price": "net_price", "Net Value": "net_value", "Incoterms": "incoterms",
}
SHIP_FIELDS = {
    "Shipment #": "shipment_number", "PROJECT": "project", "PO#": "po_number",
    "RTS Date": "rts_date", "ETA": "eta", "Delivery Date": "delivery_date",
    "Delivery Time": "delivery_time", "Status": "status", "Category": "category",
    "Supplier": "supplier", "Part Description": "part_description", "# Pcs": "num_pieces",
    "# Loads": "num_loads", "Truck Type": "truck_type", "Storage Loc": "storage_location",
    "Ship from": "ship_from", "Ship to": "ship_to", "Shipper": "shipper",
    "Shipment By (RPS/Supplier)": "shipment_by", "NCR/OSD (X)": "ncr_osd",
    "Rcvng Pics": "receiving_pics", "det pk list": "detailed_packing_list",
    "Progress Notes": "progress_notes", "Special Receiving Instructions": "special_receiving_instructions",
}
DATES = {"created_on", "delivery_date_from", "rts_date", "eta", "delivery_date"}
NUMBERS = {"ordered_quantity", "base_net_price_base_quantity_unit", "net_price", "net_value", "num_pieces", "num_loads"}
QUANTITIES = {"ordered_quantity", "num_pieces", "num_loads"}
INTEGERS = {"num_pieces", "num_loads"}


def normalize(value, field):
    if value is None or (isinstance(value, str) and not value.strip()):
        return None
    if field in NUMBERS:
        if isinstance(value, bool):
            raise ValueError("Expected a number")
        number = float(value)
        if not math.isfinite(number) or (field in QUANTITIES and number < 0):
            raise ValueError("Expected a finite, non-negative quantity")
        if field in INTEGERS:
            if not number.is_integer():
                raise ValueError("Expected a whole number")
            return int(number)
        return number
    if field in DATES or field == "item_last_change_date_time":
        if isinstance(value, str):
            try:
                value = datetime.fromisoformat(value.strip())
            except ValueError:
                for pattern in ("%m/%d/%Y", "%m/%d/%y"):
                    try:
                        value = datetime.strptime(value.strip(), pattern)
                        break
                    except ValueError:
                        pass
                else:
                    raise ValueError("Expected one valid ISO or US calendar date")
        if not isinstance(value, (datetime, date)):
            raise ValueError("Expected a calendar date")
        if field in DATES:
            return (value.date() if isinstance(value, datetime) else value).isoformat()
        return value.isoformat()
    if field == "ncr_osd":
        text = str(value).strip().lower()
        if text not in {"x", "true", "yes", "1", "false", "no", "0"}:
            raise ValueError("Expected X, yes/no, or true/false")
        return text in {"x", "true", "yes", "1"}
    if isinstance(value, time):
        return value.isoformat()
    if isinstance(value, (float, int)) and not isinstance(value, bool):
        if not math.isfinite(value):
            raise ValueError("Invalid number")
        return str(int(value)) if float(value).is_integer() else str(value)
    return str(value).strip()


def read_sheet(workbook, name, mapping, required, identifiers):
    if name not in workbook.sheetnames:
        raise ValueError(f"Missing sheet: {name}")
    rows = workbook[name].iter_rows(values_only=True)
    raw_headers = next(rows, ())
    headers, seen_headers = [], {}
    for header in raw_headers:
        header = str(header).strip() if header is not None else ""
        count = seen_headers.get(header, 0)
        seen_headers[header] = count + 1
        # The source workbook intentionally has two Product columns.
        if count and header:
            if header == "Product" and count == 1:
                header = "Product.1"
            elif header in mapping:
                raise ValueError(f"{name}: duplicate header {header}")
        headers.append(header)
    if len([h for h in headers if h in mapping]) != len(set(h for h in headers if h in mapping)):
        raise ValueError(f"{name}: duplicate mapped headers")
    missing = set(required) - set(headers)
    if missing:
        raise ValueError(f"{name}: missing headers: {', '.join(sorted(missing))}")
    records, seen = [], {}
    for row_number, values in enumerate(rows, 2):
        source = dict(zip(headers, values))
        if all(source.get(header) is None or str(source[header]).strip() == "" for header in mapping):
            continue
        record = {}
        for header, field in mapping.items():
            if header not in source:
                continue
            try:
                record[field] = normalize(source[header], field)
            except (ValueError, TypeError, OverflowError) as error:
                raise ValueError(f"{name} row {row_number}, {header}: {error}") from None
        key = tuple(record.get(field) for field in identifiers)
        if not all(key):
            raise ValueError(f"{name} row {row_number}: missing record identifier")
        if key in seen:
            raise ValueError(f"{name}: duplicate identifier in rows {seen[key]} and {row_number}")
        seen[key] = row_number
        records.append(record)
        if len(records) > 10000:
            raise ValueError(f"{name}: exceeds the 10,000-record import limit")
    if not records:
        raise ValueError(f"{name}: empty sheet; existing data has not been changed")
    return records


def read_workbook(path):
    workbook = load_workbook(path, read_only=True, data_only=True)
    formulas = None
    try:
        # Reading only cached values would turn unevaluated formulas into nulls.
        # Inspect cell types separately so neither those nor Excel errors replace
        # known-good database fields. Supplementary unmapped columns are ignored.
        formulas = load_workbook(path, read_only=True, data_only=False)
        for name, mapping in [("PO Parts Log", PO_FIELDS), ("Shipment Log", SHIP_FIELDS)]:
            if name not in formulas.sheetnames:
                continue
            raw_rows = formulas[name].iter_rows()
            cached_rows = workbook[name].iter_rows()
            headers = [str(cell.value).strip() if cell.value is not None else "" for cell in next(raw_rows, ())]
            next(cached_rows, ())
            for raw_row, cached_row in zip(raw_rows, cached_rows):
                for header, raw_cell, cached_cell in zip(headers, raw_row, cached_row):
                    if header not in mapping:
                        continue
                    if raw_cell.data_type == 'e' or cached_cell.data_type == 'e':
                        raise ValueError(f"{name} {raw_cell.coordinate}: correct the spreadsheet error before importing")
                    if raw_cell.data_type == 'f' and cached_cell.value is None:
                        raise ValueError(f"{name} {raw_cell.coordinate}: recalculate and save the workbook before importing")
        purchase_orders = read_sheet(workbook, "PO Parts Log", PO_FIELDS,
                                    ["Purchase Order ID", "Purchase Order Item", "Item Description"],
                                    ["purchase_order_id", "purchase_order_item"])
        shipments = read_sheet(workbook, "Shipment Log", SHIP_FIELDS, ["Shipment #"], ["shipment_number"])
        return purchase_orders, shipments
    finally:
        workbook.close()
        if formulas is not None:
            formulas.close()


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--file", type=Path, default=Path(os.environ.get("PO_SHIPMENT_EXCEL_FILE", Path(__file__).with_name("PO & Shipment Log.xlsx"))))
    parser.add_argument("--project-id", required=True, help="Destination project UUID; never inferred from workbook labels")
    parser.add_argument("--apply", action="store_true", help="Import validated sheets using the server-only credential")
    args = parser.parse_args(argv)
    try:
        project_id = str(UUID(args.project_id))
        purchase_orders, shipments = read_workbook(args.file)
        print(f"Validated {len(purchase_orders)} PO lines and {len(shipments)} shipments for project {project_id}.")
        if not args.apply:
            print("Validation only. No database changes. Use --apply for an intended import.")
            return 0
        url = os.environ.get("SUPABASE_URL", "").rstrip("/")
        credential = os.environ.get("SUPABASE_SERVICE_ROLE_KEY", "")
        parsed = urlsplit(url)
        if parsed.scheme != "https" or not parsed.hostname or parsed.username or parsed.password or parsed.query or parsed.fragment or parsed.path:
            raise ValueError("SUPABASE_URL must be an HTTPS origin")
        if not credential:
            raise ValueError("SUPABASE_SERVICE_ROLE_KEY is required; an anonymous key cannot import")
        response = requests.post(
            f"{url}/rest/v1/rpc/import_po_shipment_snapshot",
            headers={"apikey": credential, "Authorization": f"Bearer {credential}", "Content-Type": "application/json"},
            json={"p_project_id": project_id, "p_purchase_orders": purchase_orders, "p_shipments": shipments},
            timeout=(10, 120), allow_redirects=False,
        )
        if response.status_code != 200:
            # Do not echo arbitrary remote response bodies or credentials into logs.
            raise ValueError(f"Import rejected (HTTP {response.status_code}); check the server import error. No partial import was committed.")
        result = response.json()
        if not isinstance(result, dict) or result.get("purchase_orders") != len(purchase_orders) or result.get("shipments") != len(shipments) or not result.get("synced_at"):
            raise ValueError("Unexpected server acknowledgement. Verify import state before retrying.")
        print(f"Import committed: {result['purchase_orders']} PO lines, {result['shipments']} shipments.")
        return 0
    except requests.RequestException:
        print("Import response unavailable. The transaction may have committed; verify before retrying. Retrying preserves record IDs.", file=sys.stderr)
        return 1
    except (ValueError, OSError, KeyError, BadZipFile) as error:
        print(f"Import stopped: {error}", file=sys.stderr)
        return 1


if __name__ == "__main__":
    from dotenv import load_dotenv
    load_dotenv()
    sys.exit(main())
